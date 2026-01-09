"""Tests for file readers module."""

import pytest
from pathlib import Path
from course_scribe.core.readers import (
    read_file,
    FileReadError,
    get_supported_formats,
    _read_text,
)


class TestGetSupportedFormats:
    def test_returns_list_of_formats(self):
        formats = get_supported_formats()
        assert isinstance(formats, list)
        assert ".txt" in formats
        assert ".pdf" in formats
        assert ".docx" in formats
        assert ".pptx" in formats


class TestReadText:
    def test_reads_txt_file(self, tmp_path):
        test_file = tmp_path / "test.txt"
        test_file.write_text("Hello, World!", encoding="utf-8")

        content = read_file(test_file)
        assert content == "Hello, World!"

    def test_reads_md_file(self, tmp_path):
        test_file = tmp_path / "test.md"
        test_file.write_text("# Heading\n\nContent", encoding="utf-8")

        content = read_file(test_file)
        assert "# Heading" in content
        assert "Content" in content

    def test_reads_utf8_japanese(self, tmp_path):
        test_file = tmp_path / "test.txt"
        test_file.write_text("日本語テスト", encoding="utf-8")

        content = read_file(test_file)
        assert content == "日本語テスト"


class TestUnsupportedFormat:
    def test_raises_error_for_unsupported_format(self, tmp_path):
        test_file = tmp_path / "test.xyz"
        test_file.write_text("content")

        with pytest.raises(FileReadError) as exc_info:
            read_file(test_file)

        assert "Unsupported file format" in str(exc_info.value)
        assert ".xyz" in str(exc_info.value)


class TestReadFileNotFound:
    def test_raises_error_for_missing_file(self, tmp_path):
        missing_file = tmp_path / "nonexistent.txt"

        with pytest.raises(FileReadError):
            read_file(missing_file)


# Note: PDF, DOCX, PPTX tests require actual test files
# These are integration tests that should be run with sample files

class TestDocxReader:
    """Tests for Word document reading (requires python-docx)."""

    @pytest.fixture
    def sample_docx(self, tmp_path):
        """Create a simple test DOCX file."""
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        doc = Document()
        doc.add_heading("Test Document", 0)
        doc.add_paragraph("This is a test paragraph.")
        doc.add_paragraph("日本語テスト")

        filepath = tmp_path / "test.docx"
        doc.save(filepath)
        return filepath

    def test_reads_docx_paragraphs(self, sample_docx):
        content = read_file(sample_docx)
        assert "Test Document" in content
        assert "test paragraph" in content

    def test_reads_docx_japanese(self, sample_docx):
        content = read_file(sample_docx)
        assert "日本語テスト" in content


class TestPptxReader:
    """Tests for PowerPoint reading (requires python-pptx)."""

    @pytest.fixture
    def sample_pptx(self, tmp_path):
        """Create a simple test PPTX file."""
        try:
            from pptx import Presentation
            from pptx.util import Inches
        except ImportError:
            pytest.skip("python-pptx not installed")

        prs = Presentation()
        slide_layout = prs.slide_layouts[1]  # Title and Content
        slide = prs.slides.add_slide(slide_layout)

        title = slide.shapes.title
        title.text = "Test Slide Title"

        body = slide.shapes.placeholders[1]
        body.text = "Slide content here"

        filepath = tmp_path / "test.pptx"
        prs.save(filepath)
        return filepath

    def test_reads_pptx_title(self, sample_pptx):
        content = read_file(sample_pptx)
        assert "Test Slide Title" in content

    def test_reads_pptx_content(self, sample_pptx):
        content = read_file(sample_pptx)
        assert "Slide content here" in content

    def test_includes_slide_numbers(self, sample_pptx):
        content = read_file(sample_pptx)
        assert "Slide 1" in content
