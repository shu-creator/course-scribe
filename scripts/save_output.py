#!/usr/bin/env python3
"""Save exam preparation content to Markdown or Word files.

Usage:
    python scripts/save_output.py <output_path> [--format md|docx|both]

The content is read from stdin.

Examples:
    echo "# Content" | python scripts/save_output.py output/exam.md
    echo "# Content" | python scripts/save_output.py output/exam.docx
    echo "# Content" | python scripts/save_output.py output/exam --format both
"""

import argparse
import re
import sys
from pathlib import Path


def markdown_to_docx(markdown_content: str, output_path: Path) -> None:
    """Convert markdown content to Word document."""
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    style.font.name = 'Yu Gothic'
    style.font.size = Pt(11)

    lines = markdown_content.split('\n')
    i = 0

    while i < len(lines):
        line = lines[i]

        # Skip empty lines
        if not line.strip():
            i += 1
            continue

        # Headers
        if line.startswith('# '):
            p = doc.add_heading(line[2:].strip(), level=1)
            i += 1
            continue
        elif line.startswith('## '):
            p = doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue
        elif line.startswith('### '):
            p = doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue
        elif line.startswith('#### '):
            p = doc.add_heading(line[5:].strip(), level=4)
            i += 1
            continue

        # Horizontal rule
        if line.strip() in ['---', '***', '___']:
            doc.add_paragraph('─' * 50)
            i += 1
            continue

        # Bullet list
        if line.strip().startswith('- ') or line.strip().startswith('* '):
            text = line.strip()[2:]
            # Handle bold text
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            p = doc.add_paragraph(text, style='List Bullet')
            i += 1
            continue

        # Numbered list
        match = re.match(r'^(\d+)\.\s+(.+)$', line.strip())
        if match:
            text = match.group(2)
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            p = doc.add_paragraph(text, style='List Number')
            i += 1
            continue

        # Code block
        if line.strip().startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            if code_lines:
                p = doc.add_paragraph()
                p.style = 'No Spacing'
                run = p.add_run('\n'.join(code_lines))
                run.font.name = 'Consolas'
                run.font.size = Pt(10)
            i += 1
            continue

        # Table detection
        if '|' in line and i + 1 < len(lines) and '---' in lines[i + 1]:
            # Parse table
            table_lines = [line]
            i += 1
            while i < len(lines) and '|' in lines[i]:
                if '---' not in lines[i]:
                    table_lines.append(lines[i])
                i += 1

            if len(table_lines) >= 1:
                # Parse header
                headers = [cell.strip() for cell in table_lines[0].split('|') if cell.strip()]
                rows = []
                for tl in table_lines[1:]:
                    cells = [cell.strip() for cell in tl.split('|') if cell.strip()]
                    if cells:
                        rows.append(cells)

                if headers and rows:
                    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
                    table.style = 'Table Grid'

                    # Header row
                    for j, header in enumerate(headers):
                        if j < len(table.columns):
                            table.cell(0, j).text = header

                    # Data rows
                    for row_idx, row in enumerate(rows):
                        for col_idx, cell in enumerate(row):
                            if col_idx < len(table.columns):
                                table.cell(row_idx + 1, col_idx).text = cell

            continue

        # Regular paragraph
        text = line.strip()
        # Handle bold text
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        if text:
            p = doc.add_paragraph(text)
        i += 1

    doc.save(output_path)


def save_markdown(content: str, output_path: Path) -> None:
    """Save content as Markdown file."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(content, encoding='utf-8')


def save_docx(content: str, output_path: Path) -> None:
    """Save content as Word document."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    markdown_to_docx(content, output_path)


def main():
    parser = argparse.ArgumentParser(description='Save exam preparation content to file')
    parser.add_argument('output_path', type=Path, help='Output file path')
    parser.add_argument('--format', choices=['md', 'docx', 'both'], default=None,
                        help='Output format (default: auto-detect from extension)')

    args = parser.parse_args()

    # Read content from stdin
    content = sys.stdin.read()

    if not content.strip():
        print("Error: No content provided via stdin", file=sys.stderr)
        sys.exit(1)

    output_path = args.output_path
    output_format = args.format

    # Auto-detect format from extension if not specified
    if output_format is None:
        if output_path.suffix.lower() == '.docx':
            output_format = 'docx'
        elif output_path.suffix.lower() == '.md':
            output_format = 'md'
        else:
            output_format = 'both'

    try:
        if output_format == 'md':
            md_path = output_path if output_path.suffix == '.md' else output_path.with_suffix('.md')
            save_markdown(content, md_path)
            print(f"Saved: {md_path}")

        elif output_format == 'docx':
            docx_path = output_path if output_path.suffix == '.docx' else output_path.with_suffix('.docx')
            save_docx(content, docx_path)
            print(f"Saved: {docx_path}")

        elif output_format == 'both':
            base_path = output_path.with_suffix('')
            md_path = base_path.with_suffix('.md')
            docx_path = base_path.with_suffix('.docx')

            save_markdown(content, md_path)
            print(f"Saved: {md_path}")

            save_docx(content, docx_path)
            print(f"Saved: {docx_path}")

    except Exception as e:
        print(f"Error saving file: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
